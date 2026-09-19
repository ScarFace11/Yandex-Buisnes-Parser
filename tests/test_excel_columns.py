"""Tests for the configurable Excel export columns (web form «Выгрузка Excel»)."""
import os
import tempfile

import openpyxl

from yandex_maps_parser import state
from yandex_maps_parser.constants import CSV_FIELDS
from yandex_maps_parser.exporters import _enabled_fields, _map_field as _map_fields_for_records, save_excel

RECORDS = [
    {
        "reviewed": "",
        "name": "Тест",
        "category": "Бар",
        "phone": "+7 999 123-45-67",
        "vk": "https://vk.com/test",
        "query": "бар",
        "parsed_at": "2026-09-07",
    }
]


def _xlsx_headers(records, columns, extra=()):
    state.EXCEL_COLUMNS = columns
    try:
        path = os.path.join(tempfile.mkdtemp(), "t.xlsx")
        save_excel(records, path, extra_fields=extra)
        wb = openpyxl.load_workbook(path)
        try:
            ws = wb.active
            return [c.value for c in ws[1]], list(ws.iter_rows(values_only=True))[1]
        finally:
            wb.close()
    finally:
        state.EXCEL_COLUMNS = None


def test_enabled_fields_all_by_default():
    # The map column of the OTHER source is dropped (it would be empty).
    state.SOURCE = "yandex"
    assert _enabled_fields() == [f for f in CSV_FIELDS if f != "twogis_url"]


def test_map_column_follows_source():
    old = state.SOURCE
    try:
        state.SOURCE = "yandex"
        assert "yandex_maps_url" in _enabled_fields()
        assert "twogis_url" not in _enabled_fields()
        state.SOURCE = "2gis"
        assert "twogis_url" in _enabled_fields()
        assert "yandex_maps_url" not in _enabled_fields()
    finally:
        state.SOURCE = old


def test_website_column_is_not_exported():
    # «Сайт» held a page-wide analytics script for 2GIS runs — it is no
    # longer part of the report, only of raw files (see below).
    assert "website" not in CSV_FIELDS
    assert "website" not in _enabled_fields()


def test_raw_export_keeps_website_column():
    headers, _ = _xlsx_headers(RECORDS, None, extra=("website",))
    assert "Сайт" in headers


def test_map_column_aliased_in_user_selection():
    # Picked «2ГИС» on a Yandex run → the file must still carry a card link.
    state.SOURCE = "yandex"
    state.EXCEL_COLUMNS = {"name", "twogis_url"}
    try:
        fields = _enabled_fields()
    finally:
        state.EXCEL_COLUMNS = None
    assert fields == ["name", "yandex_maps_url"]


def test_selection_never_yields_empty_sheet():
    state.EXCEL_COLUMNS = {"2gis_only_unknown_key"}
    try:
        assert _enabled_fields()
    finally:
        state.EXCEL_COLUMNS = None


def test_map_column_follows_records_not_stale_source():
    # A 2GIS file written while state.SOURCE still says "yandex" keeps its
    # «2ГИС» column — the review-mark store keys rows by that card URL.
    state.SOURCE = "yandex"
    recs = [dict(RECORDS[0], twogis_url="https://2gis.ru/firm/123")]
    assert _map_fields_for_records(recs) == "twogis_url"
    headers, row = _xlsx_headers(recs, None)
    assert "2ГИС" in headers and "Яндекс.Карты" not in headers
    assert row[headers.index("2ГИС")] == "https://2gis.ru/firm/123"


def test_enabled_fields_subset_keeps_csv_order():
    state.EXCEL_COLUMNS = {"vk", "name", "phone"}
    try:
        assert _enabled_fields() == ["name", "phone", "vk"]
    finally:
        state.EXCEL_COLUMNS = None


def test_excel_writes_only_selected_columns():
    headers, row = _xlsx_headers(RECORDS, {"name", "phone", "vk"})
    assert headers == ["Название", "Телефон", "ВКонтакте"]
    assert row == ("Тест", "+7 999 123-45-67", "https://vk.com/test")


def test_excel_includes_all_columns_when_none_selected():
    headers, row = _xlsx_headers(RECORDS, None)
    assert len(headers) == len(_enabled_fields())
    assert "Название" in headers and "ВКонтакте" in headers
    assert row[headers.index("Название")] == "Тест"
    assert row[headers.index("ВКонтакте")] == "https://vk.com/test"