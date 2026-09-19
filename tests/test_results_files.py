"""Tests for the «Результаты» toolbox: bulk outreach, viewed marks,
file browser (RAW / PROCESSED / ARCHIVE) and the extended /results-view.

Covers the acceptance criteria of the Results-tab rework:
  * /results-view city grouping, unviewed counters and single-file view
  * /bulk/urls validation, clamping, skip_viewed and file scoping
  * /reviewed/batch limits and its effect on the counters
  * /reviewed/persist writing the «✓ Просмотрено» column atomically
  * reviewed marks flowing into /export-filtered
  * path-traversal protection for every file-taking endpoint
"""
import json
import os

import openpyxl
import pytest
from flask import Flask

from yandex_maps_parser.exporters import save_excel


# ── Fixtures ──────────────────────────────────────────────────

def _rec(name, city, url, vk=None, query="кафе"):
    r = {"name": name, "city": city, "query": query,
         "yandex_maps_url": url, "address": f"ул. {name}, 1"}
    if vk:
        r["vk"] = vk
    return r


@pytest.fixture
def client(tmp_path, monkeypatch):
    """Flask client wired to a temporary output/ tree."""
    from routes import api as api_mod

    root = tmp_path / "output"
    raw = root / "raw"
    proc = root / "processed"
    arch = root / "_archive"
    for d in (raw, proc / "excel", proc / "json", arch):
        d.mkdir(parents=True, exist_ok=True)

    monkeypatch.setattr(api_mod, "OUTPUT_DIR", str(root))
    monkeypatch.setattr(api_mod, "REVIEWED_FILE", str(root / "_reviewed.json"))
    monkeypatch.setattr(api_mod, "_raw_dir", lambda: str(raw))
    monkeypatch.setattr(api_mod, "_processed_dir", lambda: str(proc))
    monkeypatch.setattr(api_mod, "_archive_root", lambda: str(arch))
    def _today_dir():
        d = arch / "2026-01-01"
        d.mkdir(parents=True, exist_ok=True)
        return str(d)

    monkeypatch.setattr(api_mod, "_archive_today", _today_dir)
    api_mod._REC_CACHE.clear()

    app = Flask(__name__)
    app.register_blueprint(api_mod.bp)
    c = app.test_client()
    c.dirs = {"root": root, "raw": raw, "proc": proc, "arch": arch}
    return c


def _make_raw(client, name, records):
    path = client.dirs["raw"] / name
    save_excel(records, str(path))
    return path


def _make_processed(client, name, records):
    path = client.dirs["proc"] / "excel" / name
    save_excel(records, str(path))
    return path


def _mark_current_search(client, started_at, run_id="run1"):
    """Write the marker RunManager.start_process leaves at the start of a run."""
    (client.dirs["root"] / ".current_search.json").write_text(
        json.dumps({"run_id": run_id, "started_at": started_at}), encoding="utf-8")


SAMPLE = [
    _rec("Кофейня А", "Уфа", "https://yandex.ru/maps/org/a", vk="https://vk.com/a"),
    _rec("Кофейня Б", "Уфа", "https://yandex.ru/maps/org/b", vk="https://vk.com/b"),
    _rec("Бар В", "Москва", "https://yandex.ru/maps/org/v", vk="https://vk.com/v"),
    _rec("Без соцсетей", "Москва", "https://yandex.ru/maps/org/c"),
]


# ── /results-view ─────────────────────────────────────────────

class TestResultsView:
    def test_city_grouping_and_counts(self, client):
        _make_raw(client, "raw_a.xlsx", SAMPLE)
        d = client.get("/results-view?view=raw").get_json()
        assert d["count"] == 4
        cities = {c["city"]: c["count"] for c in d["cities"]}
        assert cities == {"Москва": 2, "Уфа": 2}
        # sorted by count desc, then alphabetically
        assert [c["city"] for c in d["cities"]] == ["Москва", "Уфа"]

    def test_unviewed_stats_per_social(self, client):
        _make_raw(client, "raw_a.xlsx", SAMPLE)
        d = client.get("/results-view?view=raw").get_json()
        assert d["unviewed"]["total"] == 4
        assert d["unviewed"]["by_social"]["vk"] == 3
        assert d["unviewed"]["by_social"]["telegram"] == 0

    def test_records_without_city_grouped(self, client):
        _make_raw(client, "raw_a.xlsx", [_rec("Без города", "", "https://ya.ru/1")])
        d = client.get("/results-view?view=raw").get_json()
        assert d["cities"] == [{"city": "Без города", "count": 1}]

    def test_single_file_view(self, client):
        _make_raw(client, "raw_a.xlsx", SAMPLE[:2])
        _make_raw(client, "raw_b.xlsx", SAMPLE[2:])
        d = client.get("/results-view?view=raw&file=raw/raw_b.xlsx").get_json()
        assert d["count"] == 2
        assert {r["city"] for r in d["records"]} == {"Москва"}

    def test_invalid_view_falls_back_to_raw(self, client):
        _make_raw(client, "raw_a.xlsx", SAMPLE)
        assert client.get("/results-view?view=nonsense").get_json()["view"] == "raw"

    def test_processed_view_reads_processed_dir(self, client):
        _make_raw(client, "raw_a.xlsx", SAMPLE)
        _make_processed(client, "кафе_уфа_filtered.xlsx", SAMPLE[:2])
        assert client.get("/results-view?view=processed").get_json()["count"] == 2
        # view=all merges both stages and de-duplicates by card URL
        assert client.get("/results-view?view=all").get_json()["count"] == 4

    def test_all_view_dedupes_distinct_records(self, client):
        _make_raw(client, "raw_a.xlsx", SAMPLE)
        extra = [_rec("Только в processed", "Уфа", "https://yandex.ru/maps/org/new")]
        _make_processed(client, "кафе_уфа_filtered.xlsx", SAMPLE[:2] + extra)
        assert client.get("/results-view?view=all").get_json()["count"] == 5

    def test_current_search_scope_keeps_only_the_last_search(self, client):
        """«Текущий результат» shows the last search, not every file in the folder."""
        old = _make_raw(client, "raw_old.xlsx", SAMPLE[:2])
        new = _make_raw(client, "raw_new.xlsx", SAMPLE[2:])
        os.utime(old, (1_000, 1_000))
        os.utime(new, (2_000, 2_000))
        _mark_current_search(client, started_at=1_500)
        d = client.get("/results-view?view=raw").get_json()
        assert d["scope"] == "current"
        assert d["count"] == 2
        assert {r["city"] for r in d["records"]} == {"Москва"}
        assert d["current_search"]["run_id"] == "run1"
        # Явный показ всей папки остаётся доступен
        assert client.get("/results-view?view=raw&scope=all").get_json()["count"] == 4

    def test_scope_applies_to_processed_and_to_refiltering(self, client):
        p_old = _make_processed(client, "old_filtered.xlsx", SAMPLE[:1])
        p_new = _make_processed(client, "new_filtered.xlsx", SAMPLE[1:])
        os.utime(p_old, (1_000, 1_000))
        os.utime(p_new, (2_000, 2_000))
        _mark_current_search(client, started_at=1_500)
        assert client.get("/results-view?view=processed").get_json()["count"] == 3
        # «Применить фильтры заново» пишет свежий файл — он входит в текущий поиск
        again = _make_processed(client, "new_filtered_again.xlsx", SAMPLE[1:2])
        os.utime(again, (3_000, 3_000))
        assert client.get("/results-view?view=processed").get_json()["count"] == 4

    def test_view_all_merges_both_stages_of_the_current_search(self, client):
        _make_raw(client, "raw_cafe.xlsx", SAMPLE[:2])
        _make_processed(client, "cafe_filtered.xlsx", SAMPLE[:2])
        stale_raw = _make_raw(client, "raw_stale.xlsx", SAMPLE[2:])
        os.utime(stale_raw, (1_000, 1_000))
        _mark_current_search(client, started_at=1_500)
        d = client.get("/results-view?view=all").get_json()
        assert d["count"] == 2          # только записи текущего поиска
        assert client.get("/results-view?view=all&scope=all").get_json()["count"] == 4

    def test_without_a_marker_the_whole_folder_is_shown(self, client):
        # Первый запуск .exe с готовыми файлами в output/ ничего не теряет.
        _make_raw(client, "raw_a.xlsx", SAMPLE)
        d = client.get("/results-view?view=raw").get_json()
        assert d["scope"] == "current"
        assert d["current_search"] is None
        assert d["count"] == 4

    def test_single_file_view_ignores_the_scope(self, client):
        old = _make_raw(client, "raw_old.xlsx", SAMPLE[:2])
        _make_raw(client, "raw_new.xlsx", SAMPLE[2:])
        os.utime(old, (1_000, 1_000))
        _mark_current_search(client, started_at=1_500)
        d = client.get("/results-view?view=raw&file=raw/raw_old.xlsx").get_json()
        assert d["count"] == 2

    def test_internal_and_lock_files_are_never_read(self, client, monkeypatch):
        from routes import api as api_mod
        _make_raw(client, "raw_a.xlsx", SAMPLE)
        (client.dirs["raw"] / "_seen.json").write_text("[]", encoding="utf-8")
        (client.dirs["raw"] / "~$raw_a.xlsx").write_bytes(b"lock")
        (client.dirs["proc"] / "excel" / "~$фильтр.xlsx").write_bytes(b"lock")
        seen = []
        real = api_mod._read_records_file
        monkeypatch.setattr(api_mod, "_read_records_file",
                            lambda p: (seen.append(os.path.basename(p)), real(p))[1])
        client.get("/results-view?view=all")
        assert seen == ["raw_a.xlsx"]


# ── Path safety ───────────────────────────────────────────────

class TestPathSafety:
    @pytest.mark.parametrize("bad", [
        "../config.py",
        "raw/../../config.py",
        "/etc/passwd",
        "_seen.json",
        "_reviewed.json",
        "просмотрено/старый.xlsx",
        "logs/app.log",
    ])
    def test_results_view_rejects(self, client, bad):
        assert client.get("/results-view?file=" + bad).status_code == 404

    @pytest.mark.parametrize("bad", ["../x.xlsx", "_seen.json", "processed/../_seen.json", "logs/x.log"])
    def test_files_action_rejects(self, client, bad):
        r = client.post("/files/action", json={"path": bad, "action": "delete", "confirm": True})
        assert r.status_code in (400, 404)

    def test_action_must_be_known(self, client):
        _make_raw(client, "raw_a.xlsx", SAMPLE)
        r = client.post("/files/action", json={"path": "raw/raw_a.xlsx", "action": "shred"})
        assert r.status_code == 400

    def test_missing_file_is_404(self, client):
        assert client.post("/files/action", json={"path": "raw/nope.xlsx", "action": "archive"}).status_code == 404


# ── /bulk/urls ────────────────────────────────────────────────

class TestBulkUrls:
    def test_returns_matching_urls(self, client):
        _make_raw(client, "raw_a.xlsx", SAMPLE)
        d = client.post("/bulk/urls", json={"view": "raw", "social": "vk", "count": 10}).get_json()
        assert d["total"] == 3
        assert d["returned"] == 3
        assert d["remaining"] == 0
        assert d["urls"][0]["name"]
        assert all(u["url"].startswith("https://") for u in d["urls"])

    def test_count_is_clamped(self, client):
        _make_raw(client, "raw_a.xlsx", SAMPLE)
        d = client.post("/bulk/urls", json={"view": "raw", "social": "vk", "count": 999}).get_json()
        assert d["returned"] == 3                 # only 3 available
        d2 = client.post("/bulk/urls", json={"view": "raw", "social": "vk", "count": 0}).get_json()
        assert d2["returned"] == 1                # clamped up to 1

    def test_bad_social_and_view(self, client):
        assert client.post("/bulk/urls", json={"social": "facebook"}).status_code == 400
        assert client.post("/bulk/urls", json={"view": "hack", "social": "vk"}).status_code == 400

    def test_skip_viewed_respects_marks(self, client):
        _make_raw(client, "raw_a.xlsx", SAMPLE)
        client.post("/reviewed/batch", json={"keys": ["https://yandex.ru/maps/org/a"], "reviewed": True})
        d = client.post("/bulk/urls", json={"view": "raw", "social": "vk", "count": 10,
                                            "skip_viewed": True}).get_json()
        assert d["total"] == 2
        names = {u["name"] for u in d["urls"]}
        assert "Кофейня А" not in names
        # with skip off the marked company comes back
        d2 = client.post("/bulk/urls", json={"view": "raw", "social": "vk", "count": 10,
                                             "skip_viewed": False}).get_json()
        assert d2["total"] == 3

    def test_city_filter(self, client):
        _make_raw(client, "raw_a.xlsx", SAMPLE)
        d = client.post("/bulk/urls", json={"view": "raw", "city": "Москва", "social": "vk",
                                            "count": 10}).get_json()
        assert d["total"] == 1

    def test_exclude_keys_skips_already_shown(self, client):
        _make_raw(client, "raw_a.xlsx", SAMPLE)
        first = client.post("/bulk/urls", json={"view": "raw", "social": "vk", "count": 1}).get_json()
        key = first["urls"][0]["key"]
        second = client.post("/bulk/urls", json={"view": "raw", "social": "vk", "count": 1,
                                                 "exclude_keys": [key]}).get_json()
        assert second["urls"][0]["key"] != key

    def test_file_scoping(self, client):
        _make_raw(client, "raw_a.xlsx", SAMPLE[:2])
        _make_raw(client, "raw_b.xlsx", SAMPLE[2:])
        d = client.post("/bulk/urls", json={"view": "raw", "file": "raw/raw_b.xlsx",
                                            "social": "vk", "count": 10}).get_json()
        assert {u["name"] for u in d["urls"]} == {"Бар В"}

    def test_scope_limits_the_crawl_to_the_current_search(self, client):
        old = _make_raw(client, "raw_old.xlsx", SAMPLE[:2])
        _make_raw(client, "raw_new.xlsx", SAMPLE[2:])
        os.utime(old, (1_000, 1_000))
        _mark_current_search(client, started_at=1_500)
        d = client.post("/bulk/urls", json={"view": "raw", "social": "vk",
                                            "count": 10}).get_json()
        assert {u["name"] for u in d["urls"]} == {"Бар В"}
        d_all = client.post("/bulk/urls", json={"view": "raw", "scope": "all",
                                                "social": "vk", "count": 10}).get_json()
        assert d_all["total"] == 3

    def test_file_scoping_rejects_traversal(self, client):
        assert client.post("/bulk/urls", json={"view": "raw", "file": "../../etc/passwd",
                                               "social": "vk"}).status_code == 404

    def test_non_http_urls_never_returned(self, client):
        bad = _rec("XSS", "Уфа", "https://ya.ru/x", vk="javascript:alert(1)")
        _make_raw(client, "raw_a.xlsx", [bad])
        d = client.post("/bulk/urls", json={"view": "raw", "social": "vk"}).get_json()
        assert d["total"] == 0


# ── /reviewed/batch ───────────────────────────────────────────

class TestReviewedBatch:
    def test_marks_and_unmarks(self, client):
        keys = ["https://ya.ru/1", "https://ya.ru/2"]
        d = client.post("/reviewed/batch", json={"keys": keys, "reviewed": True}).get_json()
        assert d == {"ok": True, "changed": 2}
        assert client.get("/reviewed").get_json() == {"https://ya.ru/1": True, "https://ya.ru/2": True}
        d2 = client.post("/reviewed/batch", json={"keys": keys, "reviewed": False}).get_json()
        assert d2["changed"] == 2
        assert client.get("/reviewed").get_json() == {}

    def test_idempotent_repeat(self, client):
        client.post("/reviewed/batch", json={"keys": ["k1"], "reviewed": True})
        assert client.post("/reviewed/batch", json={"keys": ["k1"], "reviewed": True}).get_json()["changed"] == 0

    def test_requires_list_and_limits(self, client):
        assert client.post("/reviewed/batch", json={"keys": "k"}).status_code == 400
        many = [f"k{i}" for i in range(501)]
        assert client.post("/reviewed/batch", json={"keys": many}).status_code == 400

    def test_blank_keys_ignored(self, client):
        d = client.post("/reviewed/batch", json={"keys": ["", None, "  ", "real"]}).get_json()
        assert d["changed"] == 1


# ── Review key ────────────────────────────────────────────────

class TestReviewKey:
    def test_prefers_yandex_then_twogis(self):
        from routes.api import _review_key
        assert _review_key({"yandex_maps_url": "y", "twogis_url": "t"}) == "y"
        assert _review_key({"yandex_maps_url": "", "twogis_url": "t"}) == "t"

    def test_fallback_composite_key(self):
        from routes.api import _review_key
        k = _review_key({"name": "Бар", "city": "Уфа", "address": "ул. Ленина, 1"})
        assert k == "n:Бар|Уфа|ул. Ленина, 1"
        assert k == _review_key({"name": "Бар", "city": "Уфа", "address": "ул. Ленина, 1"})

    def test_key_is_stable_across_views(self):
        from routes.api import _review_key
        raw = {"name": "Бар", "city": "Уфа", "address": "ул. Ленина, 1"}
        processed = dict(raw, website="x")
        assert _review_key(raw) == _review_key(processed)

    def test_empty_record_has_no_key(self):
        from routes.api import _review_key
        assert _review_key({}) == ""


# ── /files/list ───────────────────────────────────────────────

class TestFilesList:
    def test_sections_and_metadata(self, client):
        _make_raw(client, "raw_a.xlsx", SAMPLE)
        _make_processed(client, "кафе_уфа_filtered.xlsx", SAMPLE[:2])
        (client.dirs["arch"] / "2026-01-01").mkdir(parents=True, exist_ok=True)
        save_excel(SAMPLE[:1], str(client.dirs["arch"] / "2026-01-01" / "old.xlsx"))
        d = client.get("/files/list").get_json()
        assert [len(d[s]) for s in ("raw", "processed", "archive")] == [1, 1, 1]
        item = d["raw"][0]
        assert item["records"] == 4
        assert item["path"] == "raw/raw_a.xlsx"
        assert item["ext"] == "xlsx"
        assert item["size"] > 0
        assert item["modified"]
        assert item["mtime"] > 0
        assert d["archive"][0]["path"].startswith("_archive/")

    def test_newest_first(self, client):
        a = _make_raw(client, "raw_a.xlsx", SAMPLE[:1])
        b = _make_raw(client, "raw_b.xlsx", SAMPLE[:1])
        os.utime(a, (1, 1))
        os.utime(b, (2, 2))
        d = client.get("/files/list").get_json()
        assert [i["name"] for i in d["raw"]] == ["raw_b.xlsx", "raw_a.xlsx"]

    def test_unreadable_file_reports_error(self, client):
        bad = client.dirs["raw"] / "raw_broken.xlsx"
        bad.write_bytes(b"not an xlsx at all")
        d = client.get("/files/list").get_json()
        item = next(i for i in d["raw"] if i["name"] == "raw_broken.xlsx")
        assert "error" in item
        assert "records" not in item

    def test_internal_files_are_not_listed(self, client):
        (client.dirs["root"] / "_seen.json").write_text("{}", encoding="utf-8")
        (client.dirs["raw"] / "~$tmp.xlsx").write_bytes(b"x")
        d = client.get("/files/list").get_json()
        assert d["raw"] == []

    def test_json_files_counted(self, client):
        (client.dirs["proc"] / "json").mkdir(parents=True, exist_ok=True)
        (client.dirs["proc"] / "json" / "кафе.json").write_text(
            json.dumps(SAMPLE), encoding="utf-8")
        d = client.get("/files/list").get_json()
        assert d["processed"][0]["records"] == 4

    def test_record_cache_reused_and_invalidated(self, client):
        from routes import api as api_mod
        _make_raw(client, "raw_a.xlsx", SAMPLE)
        client.get("/files/list")
        assert api_mod._REC_CACHE
        client.post("/files/action", json={"path": "raw/raw_a.xlsx", "action": "archive"})
        assert api_mod._REC_CACHE == {}


# ── /files/action ─────────────────────────────────────────────

class TestFilesAction:
    def test_archive_moves_file(self, client):
        _make_raw(client, "raw_a.xlsx", SAMPLE)
        d = client.post("/files/action", json={"path": "raw/raw_a.xlsx", "action": "archive"}).get_json()
        assert d["ok"] is True
        assert (client.dirs["arch"] / "2026-01-01" / "raw_a.xlsx").is_file()
        assert not (client.dirs["raw"] / "raw_a.xlsx").exists()
        assert d["moved_to"].startswith("_archive/")

    def test_archive_renames_on_collision(self, client):
        _make_raw(client, "raw_a.xlsx", SAMPLE)
        target = client.dirs["arch"] / "2026-01-01"
        target.mkdir(parents=True, exist_ok=True)
        (target / "raw_a.xlsx").write_bytes(b"old")
        client.post("/files/action", json={"path": "raw/raw_a.xlsx", "action": "archive"})
        assert (target / "raw_a_1.xlsx").is_file()

    def test_cannot_archive_twice(self, client):
        _make_raw(client, "raw_a.xlsx", SAMPLE)
        client.post("/files/action", json={"path": "raw/raw_a.xlsx", "action": "archive"})
        r = client.post("/files/action", json={"path": "_archive/2026-01-01/raw_a.xlsx", "action": "archive"})
        assert r.status_code == 400

    def test_delete_needs_confirmation(self, client):
        _make_raw(client, "raw_a.xlsx", SAMPLE)
        r = client.post("/files/action", json={"path": "raw/raw_a.xlsx", "action": "delete"})
        assert r.status_code == 400
        assert (client.dirs["raw"] / "raw_a.xlsx").exists()

    def test_delete_with_confirmation(self, client):
        _make_raw(client, "raw_a.xlsx", SAMPLE)
        d = client.post("/files/action", json={"path": "raw/raw_a.xlsx", "action": "delete",
                                               "confirm": True}).get_json()
        assert d["ok"] is True
        assert not (client.dirs["raw"] / "raw_a.xlsx").exists()

    def test_processed_files_can_be_managed(self, client):
        _make_processed(client, "кафе_уфа_filtered.xlsx", SAMPLE)
        d = client.post("/files/action", json={"path": "processed/excel/кафе_уфа_filtered.xlsx",
                                               "action": "archive"}).get_json()
        assert d["ok"] is True
        assert (client.dirs["arch"] / "2026-01-01" / "кафе_уфа_filtered.xlsx").is_file()


class TestRestoreFromArchive:
    def _archive(self, client, rel):
        return client.post("/files/action", json={"path": rel, "action": "archive"}).get_json()["moved_to"]

    def test_restore_puts_the_file_back(self, client):
        path = _make_raw(client, "raw_a.xlsx", SAMPLE)
        archived = self._archive(client, "raw/raw_a.xlsx")
        d = client.post("/files/action", json={"path": archived, "action": "restore"}).get_json()
        assert d == {"ok": True, "restored_to": "raw/raw_a.xlsx"}
        assert path.is_file()
        assert not (client.dirs["arch"] / "2026-01-01" / "raw_a.xlsx").exists()

    def test_restore_keeps_the_original_processed_folder(self, client):
        _make_processed(client, "кафе_уфа_filtered.xlsx", SAMPLE)
        archived = self._archive(client, "processed/excel/кафе_уфа_filtered.xlsx")
        client.post("/files/action", json={"path": archived, "action": "restore"})
        assert (client.dirs["proc"] / "excel" / "кафе_уфа_filtered.xlsx").is_file()

    def test_restore_without_index_uses_the_naming_convention(self, client):
        day = client.dirs["arch"] / "2026-01-01"
        day.mkdir(parents=True, exist_ok=True)
        save_excel(SAMPLE[:1], str(day / "raw_legacy.xlsx"))
        d = client.post("/files/action", json={"path": "_archive/2026-01-01/raw_legacy.xlsx",
                                                "action": "restore"}).get_json()
        assert d["restored_to"] == "raw/raw_legacy.xlsx"
        assert (client.dirs["raw"] / "raw_legacy.xlsx").is_file()

    def test_restore_never_overwrites_an_existing_file(self, client):
        _make_raw(client, "raw_a.xlsx", SAMPLE)
        archived = self._archive(client, "raw/raw_a.xlsx")
        _make_raw(client, "raw_a.xlsx", SAMPLE)          # the original name comes back
        d = client.post("/files/action", json={"path": archived, "action": "restore"}).get_json()
        assert d["restored_to"] == "raw/raw_a_1.xlsx"
        assert len(list(client.dirs["raw"].glob("raw_a*.xlsx"))) == 2

    def test_restore_rejects_files_outside_the_archive(self, client):
        _make_raw(client, "raw_a.xlsx", SAMPLE)
        r = client.post("/files/action", json={"path": "raw/raw_a.xlsx", "action": "restore"})
        assert r.status_code == 400

    def test_index_is_not_listed_as_a_file(self, client):
        _make_raw(client, "raw_a.xlsx", SAMPLE)
        self._archive(client, "raw/raw_a.xlsx")          # creates _archive/_index.json
        names = [i["name"] for i in client.get("/files/list").get_json()["archive"]]
        assert "_index.json" not in names
        assert "raw_a.xlsx" in names


# ── /reviewed/persist (+ export) ──────────────────────────────

class TestPersistReviewed:
    def _read_reviewed_column(self, path):
        wb = openpyxl.load_workbook(path)
        ws = wb.active
        header = [str(c.value or "").strip() for c in next(ws.iter_rows(min_row=1, max_row=1))]
        col = header.index("✓ Просмотрено") + 1
        vals = [ws.cell(row=r, column=col).value for r in range(2, ws.max_row + 1)]
        wb.close()
        return vals

    def test_writes_marks_into_xlsx(self, client):
        path = _make_raw(client, "raw_a.xlsx", SAMPLE)
        client.post("/reviewed/batch", json={"keys": ["https://yandex.ru/maps/org/a"], "reviewed": True})
        d = client.post("/reviewed/persist", json={"view": "raw"}).get_json()
        assert d["ok"] and d["files"] == 1 and d["updated"] == 1 and not d["errors"]
        assert self._read_reviewed_column(path) == [True, False, False, False]

    def test_idempotent_and_no_tmp_left_behind(self, client):
        path = _make_raw(client, "raw_a.xlsx", SAMPLE)
        client.post("/reviewed/batch", json={"keys": ["https://yandex.ru/maps/org/v"], "reviewed": True})
        client.post("/reviewed/persist", json={"view": "raw"})
        d2 = client.post("/reviewed/persist", json={"view": "raw"}).get_json()
        assert d2["updated"] == 0
        assert list(client.dirs["raw"].glob("*.tmp")) == []
        assert self._read_reviewed_column(path) == [False, False, True, False]

    def test_unmarking_is_written_too(self, client):
        path = _make_raw(client, "raw_a.xlsx", SAMPLE)
        client.post("/reviewed/batch", json={"keys": ["https://yandex.ru/maps/org/b"], "reviewed": True})
        client.post("/reviewed/persist", json={"view": "raw"})
        client.post("/reviewed/batch", json={"keys": ["https://yandex.ru/maps/org/b"], "reviewed": False})
        client.post("/reviewed/persist", json={"view": "raw"})
        assert self._read_reviewed_column(path) == [False, False, False, False]

    def test_scope_writes_only_the_current_search_files(self, client):
        old = _make_raw(client, "raw_old.xlsx", SAMPLE[:2])
        new = _make_raw(client, "raw_new.xlsx", SAMPLE[2:])
        os.utime(old, (1_000, 1_000))
        _mark_current_search(client, started_at=1_500)
        client.post("/reviewed/batch", json={"keys": ["https://yandex.ru/maps/org/v"],
                                             "reviewed": True})
        d = client.post("/reviewed/persist", json={"view": "raw"}).get_json()
        assert d["files"] == 1 and d["updated"] == 1
        assert self._read_reviewed_column(new) == [True, False]
        assert not any(self._read_reviewed_column(old)), "старый файл не тронут"
        d_all = client.post("/reviewed/persist", json={"view": "raw", "scope": "all"}).get_json()
        assert d_all["files"] == 2

    def test_skips_files_without_the_column(self, client):
        path = client.dirs["raw"] / "raw_plain.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["Название", "Город"])
        ws.append(["Без колонки", "Уфа"])
        wb.save(path)
        wb.close()
        d = client.post("/reviewed/persist", json={"view": "raw"}).get_json()
        assert d["skipped"] == 1 and d["updated"] == 0

    def test_accepts_the_beacon_body_sent_as_text_plain(self, client):
        # navigator.sendBeacon при уходе со страницы может прийти без
        # JSON-заголовка — отметки всё равно должны попасть в нужный вид.
        path = _make_raw(client, "raw_a.xlsx", SAMPLE)
        client.post("/reviewed/batch", json={"keys": ["https://yandex.ru/maps/org/a"],
                                             "reviewed": True})
        d = client.post("/reviewed/persist",
                        data=json.dumps({"view": "raw", "scope": "current"}),
                        content_type="text/plain").get_json()
        assert d["ok"] and d["updated"] == 1
        assert self._read_reviewed_column(path) == [True, False, False, False]

    def test_an_empty_beacon_body_falls_back_to_the_default_view(self, client):
        _make_raw(client, "raw_a.xlsx", SAMPLE)
        r = client.post("/reviewed/persist", data="", content_type="text/plain")
        assert r.status_code == 200 and r.get_json()["ok"]

    def test_key_fallback_marks_cardless_rows(self, client):
        rec = {"name": "Без карточки", "city": "Уфа", "address": "ул. Мира, 5"}
        path = _make_raw(client, "raw_a.xlsx", [rec])
        client.post("/reviewed/batch", json={"keys": ["n:Без карточки|Уфа|ул. Мира, 5"], "reviewed": True})
        client.post("/reviewed/persist", json={"view": "raw"})
        assert self._read_reviewed_column(path) == [True]


class TestExportFiltered:
    def test_reviewed_column_reflects_marks(self, client):
        _make_raw(client, "raw_a.xlsx", SAMPLE)
        client.post("/reviewed/batch", json={"keys": ["https://yandex.ru/maps/org/a"], "reviewed": True})
        rows = SAMPLE
        r = client.post("/export-filtered", json={"rows": rows, "format": "csv"})
        assert r.status_code == 200
        text = r.data.decode("utf-8-sig")
        lines = [ln for ln in text.splitlines() if ln.strip()]
        assert "✓ Просмотрено" in lines[0]
        # first sample row is marked, the others are not
        assert lines[1].split(",")[1] == "True"
        assert lines[2].split(",")[1] == "False"

    def test_xlsx_export_has_reviewed_column(self, client):
        _make_raw(client, "raw_a.xlsx", SAMPLE)
        client.post("/reviewed/batch", json={"keys": ["https://yandex.ru/maps/org/a"], "reviewed": True})
        r = client.post("/export-filtered", json={"rows": SAMPLE, "format": "xlsx"})
        wb = openpyxl.load_workbook(__import__("io").BytesIO(r.data))
        ws = wb.active
        header = [str(c.value or "").strip() for c in next(ws.iter_rows(min_row=1, max_row=1))]
        col = header.index("✓ Просмотрено") + 1
        assert ws.cell(row=2, column=col).value is True
        assert ws.cell(row=3, column=col).value is False
        wb.close()


# ── Helpers ───────────────────────────────────────────────────

class TestHelpers:
    def test_safe_output_file_accepts_real_files(self, client):
        from routes.api import _safe_output_file
        _make_raw(client, "raw_a.xlsx", SAMPLE)
        assert _safe_output_file("raw/raw_a.xlsx")
        assert _safe_output_file("raw\\raw_a.xlsx")   # windows separators normalised
        assert _safe_output_file("raw/nope.xlsx") is None
        assert _safe_output_file("") is None

    def test_count_records_caches_by_mtime(self, client, monkeypatch):
        from routes import api as api_mod
        path = _make_raw(client, "raw_a.xlsx", SAMPLE)
        assert api_mod._count_records(str(path)) == 4
        calls = {"n": 0}
        real = api_mod._read_records_file

        def counting(p):
            calls["n"] += 1
            return real(p)

        monkeypatch.setattr(api_mod, "_read_records_file", counting)
        assert api_mod._count_records(str(path)) == 4
        assert calls["n"] == 0                      # served from the cache

    def test_cities_of_sorted_by_count_then_name(self):
        from routes.api import _cities_of
        recs = [{"city": "А"}, {"city": "Б"}, {"city": "Б"}, {"city": ""}]
        # count desc; ties broken alphabetically (so «А» before «Без города»)
        assert _cities_of(recs) == [
            {"city": "Б", "count": 2},
            {"city": "А", "count": 1},
            {"city": "Без города", "count": 1},
        ]

    def test_safe_http_url(self):
        from routes.api import _safe_http_url
        assert _safe_http_url("https://vk.com/x") == "https://vk.com/x"
        assert _safe_http_url(" http://vk.com/x ") == "http://vk.com/x"
        assert _safe_http_url("javascript:alert(1)") == ""
        assert _safe_http_url("") == ""

    def test_fill_lead_scores_scores_rows_without_one(self):
        from routes.api import _fill_lead_scores
        recs = [{"name": "A", "category": "кафе"},
                {"name": "B", "lead_score": 77}]
        out = _fill_lead_scores(recs)
        assert out[0]["lead_score"] == 30            # нет сайта +30
        assert out[1]["lead_score"] == 77            # stored score wins

    def test_view_records_carry_lead_scores(self, client):
        """A file written before scoring existed still renders scores."""
        _make_raw(client, "raw_old.xlsx", SAMPLE)     # rows have no lead_score
        data = client.get("/results-view?view=raw&scope=all").get_json()
        scores = [r.get("lead_score") for r in data["records"]]
        assert scores and all(s not in (None, "") for s in scores)

    def test_single_file_view_scores_too(self, client):
        _make_raw(client, "raw_old.xlsx", SAMPLE)
        data = client.get("/results-view?view=raw&file=raw/raw_old.xlsx").get_json()
        assert all(r.get("lead_score") not in (None, "") for r in data["records"])
