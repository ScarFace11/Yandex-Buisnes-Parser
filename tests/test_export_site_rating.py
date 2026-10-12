# -*- coding: utf-8 -*-
"""Сквозная проверка экспорта: «Сайт / Таплинк» и «Рейтинг»/«Отзывов».

Берём пару тестовых записей — с сайтом и без — и прогоняем их через все
точки выгрузки, читая файл обратно:

* полные файлы прогона (`save_csv` / `save_excel`) — обычный сайт
  организации и taplink-агрегатор уезжают в ОДНУ колонку «Сайт / Таплинк»
  (ключ отчёта остаётся `aggregator_url`), рейтинг и число отзывов с
  Яндекс.Карт — в «Рейтинг» и «Отзывов»;
* выгрузка выборки из таблицы (`/export-filtered`) — те же колонки в CSV
  и XLSX.

Запись без сайта и без рейтинга обязана дать пустые ячейки, а не прочерк
или чужое значение: пустая ячейка и есть признак «сайта нет» для этапа 2.
"""
import csv
import io

import openpyxl
import pytest
from flask import Flask

from yandex_maps_parser import state
from yandex_maps_parser.exporters import (
    _records_from_xlsx,
    save_csv,
    save_excel,
)

SITE_LABEL = "Сайт / Таплинк"
RATING_LABEL = "Рейтинг"
REVIEWS_LABEL = "Отзывов"

# Обычный сайт организации: раньше в отчёте его не было вообще.
WITH_SITE = {
    "name": "Кафе «Сайт»",
    "city": "Казань",
    "query": "кафе",
    "address": "ул. Баумана, 1",
    "phone": "+7 (843) 203-38-80",
    "website": "http://cafe-example.ru",
    "rating": "4,8",
    "reviews_count": "327",
}
# Агрегатор вместо сайта: та же колонка, значение — taplink.
WITH_TAPLINK = {
    **WITH_SITE,
    "name": "Салон «Таплинк»",
    "website": "http://taplink.cc/salon",
}
# Ни сайта, ни рейтинга — все три колонки обязаны остаться пустыми.
NO_SITE = {
    "name": "Кофейня без сайта",
    "city": "Казань",
    "query": "кофе",
    "address": "ул. Профсоюзная, 2",
    "phone": "+7 (843) 000-00-00",
}


@pytest.fixture(autouse=True)
def full_column_set(monkeypatch):
    """Пишем файлы с полным набором колонок, как в обычном прогоне.

    `state.EXCEL_COLUMNS` — модульный синглтон, который выставляет веб-форма
    «Настройка Excel-выгрузки»; тесты друг друга им не перетирают.
    """
    monkeypatch.setattr(state, "EXCEL_COLUMNS", None, raising=False)


def _csv_rows(path):
    """Полные CSV-файлы пишутся с английскими ключами (CSV_FIELDS),
    а русские заголовки — только у Excel («Сайт / Таплинк»)."""
    with open(path, encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f, delimiter=";"))


def _xlsx_headers(path):
    wb = openpyxl.load_workbook(path, read_only=True)
    try:
        ws = wb.active
        return [str(c or "").strip()
                for c in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
    finally:
        wb.close()


# ── Полные файлы прогона ──────────────────────────────────────

class TestFullReportExport:
    def test_csv_roundtrip_site_and_rating(self, tmp_path):
        path = tmp_path / "report.csv"
        save_csv([WITH_SITE, NO_SITE], str(path), append=False)
        rows = _csv_rows(path)
        assert len(rows) == 2
        site_row, bare_row = rows
        # Отдельного website в отчёте нет — обычный сайт уехал в aggregator_url
        # (единая колонка «Сайт / Таплинк»).
        assert "website" not in site_row
        assert site_row["aggregator_url"] == "https://cafe-example.ru"
        assert site_row["rating"] == "4,8"
        assert site_row["reviews_count"] == "327"
        # Без сайта и без рейтинга — пусто, а не «—».
        assert bare_row["aggregator_url"] == ""
        assert bare_row["rating"] == ""
        assert bare_row["reviews_count"] == ""

    def test_xlsx_labels_and_values(self, tmp_path):
        path = tmp_path / "report.xlsx"
        save_excel([WITH_TAPLINK, WITH_SITE, NO_SITE], str(path))
        headers = _xlsx_headers(path)
        for label in (SITE_LABEL, RATING_LABEL, REVIEWS_LABEL):
            assert label in headers, f"нет колонки «{label}»: {headers}"
        # Русской (пользовательской) колонки website в отчёте нет.
        assert "Сайт" not in headers

        by_name = {r["name"]: r for r in _records_from_xlsx(str(path))}
        assert by_name["Салон «Таплинк»"]["aggregator_url"] == "https://taplink.cc/salon"
        assert by_name["Кафе «Сайт»"]["aggregator_url"] == "https://cafe-example.ru"
        assert by_name["Кафе «Сайт»"]["rating"] == "4,8"
        assert by_name["Кафе «Сайт»"]["reviews_count"] == "327"
        # Пустые ячейки читаются обратно как отсутствие данных (не «0» и
        # не чужая ссылка) — по этому признаку этап 2 ищет «без сайта».
        assert not by_name["Кофейня без сайта"].get("aggregator_url")
        assert not by_name["Кофейня без сайта"].get("rating")
        assert not by_name["Кофейня без сайта"].get("reviews_count")

    def test_existing_taplink_is_not_overwritten_by_the_website(self, tmp_path):
        # Ссылка-агрегатор уже есть → обычный сайт остаётся в `website`
        # (нужен этапу 2), а колонка показывает агрегатор.
        both = {**WITH_SITE, "aggregator_url": "https://taplink.cc/main"}
        path = tmp_path / "report.csv"
        save_csv([both], str(path), append=False)
        rows = _csv_rows(path)
        assert rows[0]["aggregator_url"] == "https://taplink.cc/main"


# ── Выгрузка выборки из таблицы (/export-filtered) ────────────

@pytest.fixture
def client(tmp_path, monkeypatch):
    from routes import api as api_mod

    monkeypatch.setattr(api_mod, "REVIEWED_FILE",
                        str(tmp_path / "_reviewed.json"))
    app = Flask(__name__)
    app.register_blueprint(api_mod.bp)
    return app.test_client()


class TestSelectionExport:
    """Кнопки «Выборка: Excel / CSV» над таблицей результатов."""

    def _post(self, client, rows, fmt):
        resp = client.post("/export-filtered", json={"rows": rows, "format": fmt})
        assert resp.status_code == 200
        return resp

    def test_csv_has_site_and_rating_columns(self, client):
        resp = self._post(client, [WITH_SITE, NO_SITE], "csv")
        text = resp.data.decode("utf-8-sig")
        rows = list(csv.DictReader(io.StringIO(text)))
        assert len(rows) == 2
        assert rows[0][SITE_LABEL] == "https://cafe-example.ru"
        assert rows[0][RATING_LABEL] == "4,8"
        assert rows[0][REVIEWS_LABEL] == "327"
        assert rows[1][SITE_LABEL] == ""
        assert rows[1][RATING_LABEL] == ""

    def test_xlsx_has_site_and_rating_columns(self, client):
        resp = self._post(client, [WITH_SITE, NO_SITE], "xlsx")
        ws = openpyxl.load_workbook(io.BytesIO(resp.data)).active
        header = [str(c.value or "").strip()
                  for c in next(ws.iter_rows(min_row=1, max_row=1))]
        for label in (SITE_LABEL, RATING_LABEL, REVIEWS_LABEL):
            assert label in header, f"нет колонки «{label}»: {header}"

        def cell(row_idx, label):
            return ws.cell(row=row_idx, column=header.index(label) + 1).value

        assert cell(2, SITE_LABEL) == "https://cafe-example.ru"
        assert str(cell(2, RATING_LABEL)) == "4,8"
        assert str(cell(2, REVIEWS_LABEL)) == "327"
        assert not cell(3, SITE_LABEL)
        assert not cell(3, RATING_LABEL)
        assert not cell(3, REVIEWS_LABEL)
        ws.parent.close()

    def test_csv_column_names_match_the_table(self, client):
        # Заголовки выгрузки выборки = заголовки таблицы результатов, чтобы
        # файл читался без переименования колонок.
        resp = self._post(client, [WITH_SITE], "csv")
        header = resp.data.decode("utf-8-sig").splitlines()[0]
        for label in (SITE_LABEL, RATING_LABEL, REVIEWS_LABEL):
            assert label in header, f"нет колонки «{label}»: {header}"
        assert "Агрегатор" not in header
